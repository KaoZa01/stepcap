"""Guide export formats built from the guide on disk plus the finished step
pictures the page renders (so redaction is baked in one place, the page).

Markdown: a zip of `<title>.md` + `images/step-NN.<ext>`. Word: a .docx.
Pure Python (no Windows calls); test_exporters.py covers it.
"""
from __future__ import annotations

import re
import zipfile
from datetime import datetime
from io import BytesIO

_TH_MONTHS = ("ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.",
              "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค.")


def thai_date(iso: str) -> str:
    """'2026-10-16T10:00:00' -> '16 ต.ค. 2569', like the page's th-TH date."""
    try:
        d = datetime.fromisoformat(iso)
    except (TypeError, ValueError):
        return ""
    return f"{d.day} {_TH_MONTHS[d.month - 1]} {d.year + 543}"


def numbered(guide: dict) -> list[tuple[int | None, dict]]:
    """(step number, step); notes get None and don't count."""
    out, n = [], 0
    for s in guide.get("steps") or []:
        if s.get("kind") == "note":
            out.append((None, s))
        else:
            n += 1
            out.append((n, s))
    return out


def meta_line(guide: dict) -> str:
    S = guide.get("settings") or {}
    parts = []
    if S.get("step_count"):
        parts.append(f"{sum(1 for n, _ in numbered(guide) if n)} ขั้นตอน")
    if S.get("created") and thai_date(guide.get("created", "")):
        parts.append("สร้างเมื่อ " + thai_date(guide["created"]))
    if S.get("updated") and thai_date(guide.get("updated", "")):
        parts.append("แก้ไขล่าสุด " + thai_date(guide["updated"]))
    return " · ".join(parts)


def safe_title(title: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", title or "").strip() or "guide"


def md_escape(text: str) -> str:
    return re.sub(r"([\\`*_\[\]<>#|!])", r"\\\1", text or "").replace("\n", "  \n")


def _one_line(text: str) -> str:
    return md_escape(" ".join((text or "").split()))


def build_markdown(guide: dict, images: dict[str, tuple[str, bytes]]) -> bytes:
    """images: step id -> (ext, bytes); a step without one has no picture."""
    steps = numbered(guide)
    lines = [f"# {_one_line(guide.get('title', ''))}", ""]
    meta = meta_line(guide)
    if meta:
        lines += [meta, ""]
    if (guide.get("settings") or {}).get("toc") and any(n for n, _ in steps):
        lines += ["**สารบัญ**", ""]
        lines += [f"{n}. [{_one_line(s.get('text'))}](#step-{n})" for n, s in steps if n]
        lines.append("")
    files = {}
    for n, s in steps:
        if n is None:
            lines += [f"> {md_escape(line)}" for line in (s.get("text") or "").split("\n")]
            lines.append("")
            continue
        lines += [f'<a id="step-{n}"></a>', f"## {n}. {_one_line(s.get('text'))}", ""]
        if s.get("id") in images:
            ext, data = images[s["id"]]
            name = f"images/step-{n:02d}.{ext}"
            files[name] = data
            lines += [f"![]({name})", ""]
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(safe_title(guide.get("title", "")) + ".md", "\n".join(lines))
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------
THAI_FONT = "Leelawadee UI"


def _thai_font(style) -> None:
    """Pin a style's fonts (incl. complex script, which Thai uses) to THAI_FONT;
    theme-font attributes would win over them, so they go."""
    from docx.oxml.ns import qn
    rpr = style.element.get_or_add_rPr()
    f = rpr.find(qn("w:rFonts"))
    if f is None:
        f = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, f)
    for a in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        f.attrib.pop(qn("w:" + a), None)
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        f.set(qn("w:" + a), THAI_FONT)


def _png_bytes(ext: str, data: bytes) -> bytes:
    """Word shows PNG; a GIF becomes its first frame."""
    if ext == "png":
        return data
    from PIL import Image
    with Image.open(BytesIO(data)) as im:
        out = BytesIO()
        im.convert("RGB").save(out, "PNG")
        return out.getvalue()


def _run(p, text: str, bold=False, italic=False):
    """A run whose bold/italic also reach Thai, which Word styles as complex
    script (bCs/iCs); plain b/i would leave the Thai part regular."""
    r = p.add_run(text)
    rpr = r._r.get_or_add_rPr()
    if bold:
        r.bold = True
        rpr.get_or_add_bCs()
    if italic:
        r.italic = True
        rpr.get_or_add_iCs()
    return r


def build_docx(guide: dict, images: dict[str, tuple[str, bytes]]) -> bytes:
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Emu, Inches, Mm
    from PIL import Image
    doc = docx.Document()
    for name in ("Normal", "Title", "List Number"):
        _thai_font(doc.styles[name])
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)  # A4
    text_w = sec.page_width - sec.left_margin - sec.right_margin
    steps = numbered(guide)

    doc.add_paragraph(guide.get("title", ""), style="Title")
    meta = meta_line(guide)
    if meta:
        doc.add_paragraph(meta)
    if (guide.get("settings") or {}).get("toc") and any(n for n, _ in steps):
        _run(doc.add_paragraph(), "สารบัญ", bold=True)
        for n, s in steps:
            if n:
                doc.add_paragraph(s.get("text", ""), style="List Number")
    for n, s in steps:
        if n is None:
            p = doc.add_paragraph()
            _run(p, s.get("text") or "", italic=True)
            shd = p._p.get_or_add_pPr().makeelement(
                qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "F3F4F6"})
            p._p.get_or_add_pPr().append(shd)
            continue
        _run(doc.add_paragraph(), f"{n}. {s.get('text') or ''}", bold=True)
        if s.get("id") in images:
            png = _png_bytes(*images[s["id"]])
            with Image.open(BytesIO(png)) as im:
                w = Inches(im.width / 96)
            doc.add_picture(BytesIO(png), width=Emu(min(w, text_w)))
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
