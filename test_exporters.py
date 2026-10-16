"""Self-check for exporters.py (Markdown zip, Word). Run: python test_exporters.py"""
import sys
import zipfile
from io import BytesIO

from PIL import Image

from exporters import build_markdown, md_escape, meta_line, thai_date


def _png(color=(0, 200, 0)) -> bytes:
    b = BytesIO()
    Image.new("RGB", (4, 4), color).save(b, "PNG")
    return b.getvalue()


PNG = _png()
G = {"title": "คู่มือ: ทดสอบ", "created": "2026-10-16T10:00:00", "updated": "2026-10-16T10:00:00",
     "settings": {"toc": True, "step_count": True, "created": True, "updated": False},
     "steps": [{"id": "a", "kind": "click", "text": "คลิก *บันทึก*"},
               {"id": "n", "kind": "note", "text": "ระวัง"},
               {"id": "b", "kind": "image", "text": "ภาพ"}]}


def _zip(b: bytes) -> dict:
    with zipfile.ZipFile(BytesIO(b)) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_thai_date():
    assert thai_date("2026-10-16T10:00:00") == "16 ต.ค. 2569"
    assert thai_date("x") == ""


def test_meta_line():
    assert meta_line(G) == "2 ขั้นตอน · สร้างเมื่อ 16 ต.ค. 2569"


def test_md_layout():
    z = _zip(build_markdown(G, {"a": ("png", PNG)}))
    assert set(z) == {"คู่มือ_ ทดสอบ.md", "images/step-01.png"}, set(z)
    md = z["คู่มือ_ ทดสอบ.md"].decode("utf-8")
    assert md.startswith("# คู่มือ: ทดสอบ\n"), md
    assert "1. [คลิก \\*บันทึก\\*](#step-1)" in md and '<a id="step-1"></a>' in md, md
    assert "## 1. คลิก \\*บันทึก\\*\n" in md and "![](images/step-01.png)" in md, md
    assert "> ระวัง" in md and "## 2. ภาพ" in md, md
    assert "step-02" not in md.split("## 2. ภาพ")[1], md
    assert z["images/step-01.png"] == PNG


def test_md_escapes_text():
    assert md_escape("1. # a*b_c [x](y) <i> | `q` !") == \
        "1\\. \\# a\\*b\\_c \\[x\\](y) \\<i\\> \\| \\`q\\` \\!"
    assert md_escape("a & b ~~c~~ =") == "a \\& b \\~\\~c\\~\\~ \\="
    # line starts that would turn a note into a list, rule or code block
    assert md_escape("- a\n+ b\n2) c\n    d\n---") == "\\- a  \n\\+ b  \n2\\) c  \nd  \n\\---"


def test_docx_metadata_is_guide_only():
    from exporters import build_docx
    cp = _doc(build_docx(G, {})).core_properties
    assert cp.title == G["title"] and cp.author == "" and cp.comments == ""
    assert cp.last_modified_by == "" and cp.created.year == 2026


def test_md_no_images():
    z = _zip(build_markdown({**G, "steps": [G["steps"][1]]}, {}))
    assert list(z) == ["คู่มือ_ ทดสอบ.md"], list(z)


def test_md_keeps_gif():
    assert _zip(build_markdown(G, {"b": ("gif", b"GIF89a..")}))["images/step-02.gif"] == b"GIF89a.."


def test_md_safe_inner_name():
    assert list(_zip(build_markdown({**G, "title": ' a/b\\c?*"<>| '}, {}))) == ["a_b_c______.md"]


def _doc(b: bytes):
    import docx
    return docx.Document(BytesIO(b))


def _media(b: bytes) -> list[bytes]:
    return [v for k, v in _zip(b).items() if k.startswith("word/media/")]


def test_docx_text_and_images():
    from exporters import build_docx
    b = build_docx(G, {"a": ("png", PNG)})
    d = _doc(b)
    texts = [p.text for p in d.paragraphs]
    for t in ("คู่มือ: ทดสอบ", "2 ขั้นตอน · สร้างเมื่อ 16 ต.ค. 2569", "1. คลิก *บันทึก*", "ระวัง", "2. ภาพ"):
        assert t in texts, (t, texts)
    assert texts.count("คลิก *บันทึก*") == 1, texts  # the TOC entry
    assert len(d.inline_shapes) == 1
    assert _media(b) == [PNG]


def test_docx_thai_font():
    from docx.oxml.ns import qn
    from exporters import build_docx
    rf = _doc(build_docx(G, {})).styles["Normal"].element.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert rf.get(qn("w:cs")) == "Leelawadee UI" and rf.get(qn("w:ascii")) == "Leelawadee UI"


def test_docx_thai_bold_italic():
    # Word styles Thai as complex script: bold/italic need bCs/iCs, not just b/i
    from docx.oxml.ns import qn
    from exporters import build_docx
    runs = {r.text: r._r.find(qn("w:rPr")) for p in _doc(build_docx(G, {})).paragraphs for r in p.runs}
    for t in ("1. คลิก *บันทึก*", "สารบัญ"):
        assert runs[t].find(qn("w:bCs")) is not None, t
    assert runs["ระวัง"].find(qn("w:iCs")) is not None


def test_docx_gif_first_frame():
    from exporters import build_docx
    frames = [Image.new("RGB", (8, 8), c) for c in ((255, 0, 0), (0, 0, 255))]
    g = BytesIO()
    frames[0].save(g, "GIF", save_all=True, append_images=frames[1:])
    b = build_docx(G, {"b": ("gif", g.getvalue())})
    assert len(_doc(b).inline_shapes) == 1
    m = _media(b)
    assert len(m) == 1 and m[0].startswith(b"\x89PNG")
    assert Image.open(BytesIO(m[0])).getpixel((0, 0))[:3] == (255, 0, 0)


def test_docx_notes_only():
    from exporters import build_docx
    d = _doc(build_docx({**G, "steps": [G["steps"][1]]}, {}))
    assert len(d.inline_shapes) == 0 and "ระวัง" in [p.text for p in d.paragraphs]


def _app():
    """app.py in this process on its own temp storage (never the real one)."""
    import os
    import tempfile
    from pathlib import Path
    if "app" not in sys.modules:
        os.environ["STEPCAP_STORAGE"] = tempfile.mkdtemp(prefix="stepcap-exp-")
    import app
    assert str(app.STORAGE) == os.environ["STEPCAP_STORAGE"]
    return app


def test_export_images_binary_and_missing():
    import json
    app = _app()
    st = app.STORAGE
    (st / "g2").mkdir()
    (st / "g2" / "bad.gif").write_bytes(b"GIF89a junk")
    guide = {"id": "g2", "steps": [
        {"id": "s1", "kind": "click", "text": "a", "image": "a.png"},
        {"id": "n", "kind": "note", "text": "x"},
        {"id": "s2", "kind": "image", "text": "b", "image": "gone.gif"},
        {"id": "s3", "kind": "image", "text": "c", "image": "bad.gif"}]}
    head = json.dumps({"images": [{"id": "s1", "ext": "png", "size": len(PNG)},
                                  {"id": "s2", "ext": "gif"}, {"id": "s3", "ext": "gif"}]})
    images, missing = app._export_images("g2", guide, head.encode() + b"\n" + PNG)
    assert images == {"s1": ("png", PNG)} and missing == [2, 3], (images, missing)
    images, missing = app._export_images("g2", guide, head.encode() + b"\n" + PNG[:10])
    assert images == {} and missing == [1, 2, 3], (images, missing)  # cut-off body


def test_safe_title_windows_names():
    from exporters import safe_title
    assert safe_title("con") == "_con" and safe_title("NUL.txt") == "_NUL.txt"
    assert safe_title("a\x07b. . ") == "a_b" and safe_title("...") == "guide"


def test_export_write_error_reported():
    # Save As picked a place that can't be written: say so, don't pretend it saved
    import json
    app = _app()
    st = app.STORAGE
    (st / "g1").mkdir()
    (st / "g1" / "guide.json").write_text(json.dumps(
        {"id": "g1", "title": "t", "created": "2026-10-16T10:00:00", "steps": []}), encoding="utf-8")
    bad = str(st / "no" / "such" / "x.html")
    app.WINDOW = type("W", (), {"create_file_dialog": lambda self, *a, **k: (bad,)})()
    for fmt, body in (("html", b"<html></html>"), ("md", b'{"images": []}\n')):
        assert app.export_guide("g1", fmt, body) == {"saved": None, "error": "บันทึกไฟล์ไม่ได้"}
    assert app.download_image("g1", "1", "png", PNG) == {"saved": None, "error": "บันทึกไฟล์ไม่ได้"}


if __name__ == "__main__":
    import traceback
    sys.stdout.reconfigure(encoding="utf-8")
    fns = [v for k, v in list(globals().items()) if k.startswith("test_")]
    failed = 0
    for fn in fns:
        try:
            fn()
            print("ok  ", fn.__name__)
        except Exception:
            failed += 1
            print("FAIL", fn.__name__)
            traceback.print_exc()
    print(f"\n{len(fns) - failed}/{len(fns)} passed")
    sys.exit(1 if failed else 0)
